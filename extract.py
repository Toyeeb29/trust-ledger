"""Step 1a - Pull commitments out of completed security questionnaires.

    python extract.py samples/*.xlsx            # rule-based (offline, deterministic)
    python extract.py samples/*.xlsx --llm      # Claude classifies; falls back to rules per row on error

Handles different layouts: finds the sheet and header row by itself, then merges every answer-ish
column (Response / Answer / Notes / Implementation Description...) into one answer string.

Every answer is classified as:
  commitment  an operational promise that can drift or break (timeframes, frequencies, scope)
  roadmap     a future promise ("expected Q1 2027") - tracked with a due date
  attestation a plain current-state claim (Yes / No / short fact)
  n/a

Output: out/claims.json (everything) and out/review_queue.csv (commitments + roadmap for a human to approve).
"""
import argparse
import csv
import json
import os
import re
from pathlib import Path

from openpyxl import load_workbook

import taxonomy as tx

ROOT = Path(__file__).parent
OUT = ROOT / "out"

Q_HEADERS = ("question",)
A_HEADERS = ("response", "answer", "additional", "description", "notes", "comment", "explanation")
ID_HEADERS = ("ques num", "question id", "#", "id", "ref")


# ---- reading arbitrary questionnaire layouts --------------------------------------------------
def find_table(wb):
    """Return (sheet, header_row_idx, columns) for the sheet that looks most like a questionnaire."""
    best = None
    for ws in wb.worksheets:
        for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), 1):
            cells = [str(c).strip().lower() if c is not None else "" for c in row]
            # 'Question ID' / 'Question #' are identifiers, not the question text
            q = [i for i, c in enumerate(cells) if any(h in c for h in Q_HEADERS)
                 and not re.search(r"\b(id|num|no|#)\b", c)]
            a = [i for i, c in enumerate(cells) if any(h in c for h in A_HEADERS) and i not in q]
            if q and a:
                ident = next((i for i, c in enumerate(cells) if c in ID_HEADERS or c.startswith(ID_HEADERS)), None)
                cand = (ws, r_idx, {"q": q[0], "a": a, "id": ident})
                if best is None or ws.max_row > best[0].max_row:
                    best = cand
                break
    if not best:
        raise ValueError("No questionnaire table found (need a 'Question' column and a response column)")
    return best


def read_rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws, header_row, cols = find_table(wb)
    for r_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        question = row[cols["q"]] if cols["q"] < len(row) else None
        if not question:
            continue
        parts = [str(row[i]).strip() for i in cols["a"] if i < len(row) and row[i] not in (None, "")]
        ref = str(row[cols["id"]]) if cols["id"] is not None and row[cols["id"]] is not None else f"row {r_idx}"
        yield {"ref": ref, "sheet": ws.title, "row": r_idx, "question": str(question).strip(),
               "answer": combine(parts)}


def combine(parts):
    """'Yes' + 'Security logs are kept 1 year.' -> 'Security logs are kept 1 year.' (drop redundant Yes/No)."""
    if len(parts) > 1 and parts[0].lower() in ("yes", "no", "na", "n/a"):
        return " ".join(parts[1:])
    return " ".join(parts)


# ---- rule-based classifier --------------------------------------------------------------------
def classify_rules(question, answer):
    a = answer.strip()
    topic = tx.match_topic(question)
    base = {"topic": topic, "value": None, "value_display": None, "strength": None,
            "clock_start": None, "due": None, "method": "rules"}

    if a.lower() in tx.NA_VALUES or not a:
        return {**base, "kind": "n/a", "confidence": "high"}
    if tx.ROADMAP_RE.search(a):
        return {**base, "kind": "roadmap", "topic": roadmap_subject(question, a), "due": due_date(a),
                "confidence": "high"}
    if topic:
        kind_ = tx.TOPICS[topic]["kind"]
        value, display = tx.parse_value(kind_, a)
        return {**base, "kind": "commitment", "value": value, "value_display": display,
                "strength": "soft" if tx.HEDGE_RE.search(a) else "firm",
                "clock_start": tx.clock_start(a) if kind_ == "duration" else None,
                "confidence": "high" if value is not None else "medium"}
    return {**base, "kind": "attestation", "confidence": "high" if len(a) < 60 else "medium"}


NAMED_SUBJECTS = r"SOC 2 Type II|SOC 2|ISO 27001|ISO 42001|HITRUST|FedRAMP|PCI DSS|SIEM|DLP|EDR|SSO|MFA"


def roadmap_subject(question, answer):
    for text in (question, answer):  # the question names the thing precisely; answers often abbreviate
        m = re.search(rf"\b({NAMED_SUBJECTS})\b", text)
        if m:
            return m.group(1)
    return question[:60]


def due_date(text):
    m = re.search(r"\bQ([1-4])\s*(20\d\d)\b", text, re.I)
    if m:
        return f"{m.group(2)}-{int(m.group(1)) * 3:02d}-{[31, 30, 30, 31][int(m.group(1)) - 1]}"
    m = re.search(r"\bend of (20\d\d)\b", text, re.I)
    if m:
        return f"{m.group(1)}-12-31"
    return None


# ---- optional LLM classifier ------------------------------------------------------------------
LLM_PROMPT = """You review answers a SaaS vendor gave in a customer security questionnaire.
Classify the answer and normalise it. Reply with JSON only:
{{"kind": "commitment|roadmap|attestation|n/a",
  "topic": one of {topics} or null,
  "value_display": short normalised value e.g. "72 hours", "quarterly", "all workforce accounts", or null,
  "strength": "firm|soft" or null (soft = hedged: typically, target, best effort),
  "clock_start": what event starts the timer, or null,
  "due": "YYYY-MM-DD" for roadmap items with a date, else null}}
A commitment is an operational promise that could later be broken (timeframes, frequencies, retention, scope).
Question: {q}
Answer: {a}"""


def classify_llm(question, answer, client):
    rules = classify_rules(question, answer)
    try:
        msg = client.messages.create(
            model=os.environ.get("TRUST_LEDGER_MODEL", "claude-haiku-4-5-20251001"), max_tokens=300,
            messages=[{"role": "user", "content": LLM_PROMPT.format(topics=list(tx.TOPICS), q=question, a=answer)}])
        data = json.loads(re.search(r"\{.*\}", msg.content[0].text, re.S).group(0))
        topic = data.get("topic") if data.get("topic") in tx.TOPICS else None
        # Keep numeric values deterministic: parse them from the text with the rule parser.
        value = tx.parse_value(tx.TOPICS[topic]["kind"], answer)[0] if topic else None
        return {**rules, **{k: data.get(k) for k in ("kind", "value_display", "strength", "clock_start", "due")},
                "topic": topic, "value": value, "method": "llm", "confidence": "medium"}
    except Exception as e:  # noqa: BLE001 - any failure falls back to rules
        return {**rules, "method": f"rules (llm failed: {type(e).__name__})"}


# ---- main -------------------------------------------------------------------------------------
def customer_name(path):
    return " ".join(w.capitalize() for w in Path(path).stem.split("_")[:2])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    ap.add_argument("--llm", action="store_true", help="classify with Claude (needs ANTHROPIC_API_KEY)")
    args = ap.parse_args()

    client = None
    if args.llm:
        import anthropic  # pip install anthropic
        client = anthropic.Anthropic()

    claims = []
    for f in args.files:
        cust = customer_name(f)
        for row in read_rows(f):
            c = classify_llm(row["question"], row["answer"], client) if client else classify_rules(row["question"], row["answer"])
            claims.append({"id": f"C{len(claims) + 1:04d}", "customer": cust, "source_file": Path(f).name, **row, **c})

    OUT.mkdir(exist_ok=True)
    (OUT / "claims.json").write_text(json.dumps(claims, indent=2))

    review = [c for c in claims if c["kind"] in ("commitment", "roadmap")]
    with open(OUT / "review_queue.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["claim_id", "customer", "source_ref", "question", "answer", "kind", "topic", "value",
                    "strength", "clock_start", "due", "confidence", "approved (Y/N)", "reviewer", "notes"])
        for c in review:
            w.writerow([c["id"], c["customer"], f'{c["source_file"]} {c["ref"]}', c["question"], c["answer"], c["kind"],
                        c["topic"] or "", c["value_display"] or "", c["strength"] or "", c["clock_start"] or "",
                        c["due"] or "", c["confidence"], "", "", ""])

    by_kind = {}
    for c in claims:
        by_kind[c["kind"]] = by_kind.get(c["kind"], 0) + 1
    print(f"Read {len(claims)} answers from {len(args.files)} questionnaires: {by_kind}")
    print(f"Review queue: {len(review)} items -> out/review_queue.csv")


if __name__ == "__main__":
    main()
